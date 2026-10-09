"""Utilidades transversales: limpieza de cédulas y montos, lectura de Excel con
hipervínculos y el control del recolector de basura en hilos de trabajo.

Copiado de axioma/nucleo/utils.py sin modificar el código: solo se dejó lo que Axio usa.
La red de seguridad de los finalizadores de tkinter, que en Axioma vive en ese mismo
archivo, aquí está en axio/ui/hilos.py: importa tkinter, y nucleo/ no puede hacerlo.
"""

from contextlib import contextmanager
from openpyxl import load_workbook
import gc
import io

import pandas as pd
import re
import urllib.request


# Nombre real de la columna de recibo en la Matriz_Nube. Se llama RWS -- NO RCW, aunque
# medio código (y media empresa) la nombre así de palabra. La confusión ya costó DOS bugs
# silenciosos idénticos, en dos módulos distintos, cada uno descubierto por separado:
#   - el candado de recibo del motor de conciliación no cruzaba nada nunca;
#   - Axio nunca pintaba de verde una fila ya ingresada.
# En los dos casos la comparación era exacta contra 'RCW', el nombre no coincidía, y ni
# uno ni otro avisaba de nada. Vive aquí, en un solo sitio, justamente para que el día
# que la hoja cambie de nombre otra vez sea UNA edición y no una cacería.
ALIAS_COLUMNA_RECIBO_NUBE = ('RWS', 'RCW')


def limpiar_cedula(x):
    if pd.isna(x) or x is None: return '0'
    s = str(x).strip().upper()
    
    # 1. Rescate por si Excel lo pasó a notación científica (Ej: 1.17E+07)
    if 'E+' in s or 'E-' in s:
        try: s = str(int(float(s)))
        except: pass
        
    # 2. Cortar los "falsos decimales" de Excel al final (Ej: ",00" o ".0")
    s = re.sub(r'[,.]0+$', '', s)
    
    # 3. Eliminar todos los puntos y comas restantes (separadores de miles colombianos)
    s = s.replace('.', '').replace(',', '')
    
    # 4. Destruir cualquier letra o espacio, dejando SOLO los números puros
    s = re.sub(r'\D', '', s) 
    
    return s if s != '' else '0'

def parse_money(x):
    try:
        if pd.isna(x) or x is None: return 0.0
        if isinstance(x, (int, float)): return float(x)
        s = str(x).strip()
        if not s: return 0.0
        signo = -1.0 if '-' in s or '(' in s else 1.0
        s = re.sub(r'[^\d.,]', '', s)
        if not s: return 0.0
        if '.' in s and ',' in s:
            if s.rfind(',') > s.rfind('.'): s = s.replace('.', '').replace(',', '.')
            else: s = s.replace(',', '')
        elif ',' in s:
            partes = s.split(',')
            if len(partes) > 2 or len(partes[-1]) == 3: s = s.replace(',', '')
            else: s = s.replace(',', '.')
        elif '.' in s:
            partes = s.split('.')
            if len(partes) > 2 or len(partes[-1]) == 3: s = s.replace('.', '')
        return float(s) * signo
    except: return 0.0


def _formatear_cco(valor):
    """Normaliza un código de congregación: si es numérico, lo rellena con
    ceros a la izquierda hasta 4 dígitos (ej. 35 -> '0035'); si llegó con
    basura de digitación (salto de línea pegado, texto en vez de número), se
    limpia de espacios/saltos de línea pero se deja tal cual -- se prefiere
    mostrar el dato sucio a inventar un código que no es."""
    texto = str(valor).replace('\n', '').replace('\r', '').strip()
    texto_sin_punto = texto[:-2] if texto.endswith('.0') else texto
    if texto_sin_punto.isdigit():
        return texto_sin_punto.zfill(4)
    return texto


def _leer_workbook_openpyxl(url_o_ruta):
    """Abre un .xlsx ya sea desde un path local o una URL http(s), devolviendo
    un libro CON LOS HIPERVÍNCULOS INTACTOS -- por eso aquí NO se usa pandas:
    pd.read_excel() descarta cualquier hipervínculo de celda, y un CSV (lo que
    descargar_base_global pide para Matriz_Nube) ni siquiera puede llevarlos.

    Primero se intenta con LibroRapido (mismo resultado, ~4 veces más rápido); si
    ese lector no entiende el archivo, se abre con openpyxl como siempre."""
    if url_o_ruta.startswith('http://') or url_o_ruta.startswith('https://'):
        # 120s (no 30s) -- una línea 'bastante extensa' exportada como xlsx
        # completo desde Google Sheets puede tardar bastante más que una hoja
        # chica; con 30s se corría el riesgo de cortar la descarga a medias
        # en una línea grande y fallar en silencio.
        with urllib.request.urlopen(url_o_ruta, timeout=120) as respuesta:
            datos = respuesta.read()
        fuente = io.BytesIO(datos)
    else:
        fuente = url_o_ruta
    try:
        return LibroRapido(fuente)
    except Exception:
        if hasattr(fuente, 'seek'):
            fuente.seek(0)
        return load_workbook(fuente, data_only=True)


# ==============================================================================
# LECTOR RÁPIDO DE XLSX (solo valores e hipervínculos)
# ==============================================================================
# openpyxl arma un objeto con estilo por cada celda: en una línea de crédito son ~320.000
# celdas y casi todo el tiempo de carga (en el navegador, con WebAssembly, ~30 s de las 7
# líneas). Las líneas solo necesitan el valor y el hipervínculo de cada celda, así que este
# lector lee el XML directo y entrega lo mismo que openpyxl con data_only=True:
#   - los mismos valores (números int/float, fechas según el formato de la celda, textos,
#     errores como '#REF!', '#VALUE!' para fechas fuera de rango);
#   - las celdas combinadas vacías salvo la de arriba a la izquierda;
#   - los hipervínculos (también por rango y sobre celdas combinadas);
#   - las mismas dimensiones de iter_rows(), contando las celdas que openpyxl crea al
#     asignar hipervínculos y comentarios.
# Comprobado celda por celda contra openpyxl con las 7 líneas reales
# (pruebas/verificar_lector_xlsx.py). Si algo del archivo no se entiende, se levanta un
# error y _leer_workbook_openpyxl usa openpyxl.
_NS_XLSX = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_NS_REL_DOC = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
_NS_REL_PKG = '{http://schemas.openxmlformats.org/package/2006/relationships}'
_DIGITOS = '0123456789'


class _Enlace:
    __slots__ = ('target',)

    def __init__(self, target):
        self.target = target


class _Celda:
    __slots__ = ('value', 'hyperlink')

    def __init__(self, value=None, hyperlink=None):
        self.value = value
        self.hyperlink = hyperlink


_CELDA_VACIA = _Celda()   # compartida: nadie modifica las celdas que entrega el lector


def _columna_a_numero(letras, _cache={}):
    n = _cache.get(letras)
    if n is None:
        n = 0
        for ch in letras:
            n = n * 26 + (ord(ch) - 64)
        _cache[letras] = n
    return n


def _coordenada(ref):
    letras = ref.rstrip(_DIGITOS)
    return int(ref[len(letras):]), _columna_a_numero(letras)


def _rango(ref):
    """'B2:D5' -> (2, 2, 5, 4); 'B2' -> (2, 2, 2, 2). En mayúsculas, sin '$'."""
    ref = ref.replace('$', '').upper()
    if ':' in ref:
        a, b = ref.split(':', 1)
        (f1, c1), (f2, c2) = _coordenada(a), _coordenada(b)
        return min(f1, f2), min(c1, c2), max(f1, f2), max(c1, c2)
    f, c = _coordenada(ref)
    return f, c, f, c


def _texto_de(nodo):
    """Como openpyxl Text.content: el <t> directo más el <t> de cada <r>, sin fonética."""
    partes = []
    t = nodo.find(_NS_XLSX + 't')
    if t is not None and t.text is not None:
        partes.append(t.text)
    for r in nodo.findall(_NS_XLSX + 'r'):
        t = r.find(_NS_XLSX + 't')
        if t is not None and t.text is not None:
            partes.append(t.text)
    return ''.join(partes)


def _rels(archivo, ruta):
    """{Id: destino} de un .rels, con las rutas internas normalizadas como openpyxl."""
    import posixpath
    import xml.etree.ElementTree as ET
    carpeta = posixpath.dirname(posixpath.dirname(ruta))   # .../_rels/x.rels -> ...
    rels = {}
    if ruta not in archivo.namelist():
        return rels
    for r in ET.fromstring(archivo.read(ruta)).iter(_NS_REL_PKG + 'Relationship'):
        destino = r.get('Target')
        if r.get('TargetMode') != 'External':
            destino = destino[1:] if destino.startswith('/') else posixpath.normpath(posixpath.join(carpeta, destino))
        rels.setdefault(r.get('Id'), (destino, r.get('Type', '')))
    return rels


def _ruta_rels(ruta):
    import posixpath
    carpeta, nombre = posixpath.split(ruta)
    return posixpath.join(carpeta, '_rels', nombre + '.rels')


class HojaRapida:
    def __init__(self, libro, titulo, ruta):
        self._libro = libro
        self.title = titulo
        self._ruta = ruta
        self._celdas = None

    def _valor(self, tipo, valor, estilo):
        """El valor de una celda como lo deja openpyxl (parse_cell con data_only=True)."""
        if tipo == 'n':
            valor = float(valor) if ('.' in valor or 'E' in valor or 'e' in valor) else int(valor)
            if estilo in self._libro._formatos_fecha:
                from openpyxl.utils.datetime import from_excel
                try:
                    return from_excel(valor, self._libro._epoca, timedelta=estilo in self._libro._formatos_duracion)
                except (OverflowError, ValueError):
                    return '#VALUE!'
            return valor
        if tipo == 's':
            return self._libro._compartidos[int(valor)]
        if tipo == 'b':
            return bool(int(valor))
        if tipo == 'd':
            from openpyxl.utils.datetime import from_ISO8601
            return from_ISO8601(valor)
        if tipo in ('str', 'e'):
            return valor
        raise ValueError(f"tipo de celda desconocido: {tipo}")

    # Una celda con contenido, tal como la escriben Google Sheets y Excel: <c r="B7" s="3"
    # t="s"><v>12</v></c>, a veces con <f>…</f> antes. Las vacías (<c r="B7" s="3"/>) no
    # coinciden y la búsqueda se las salta sin pasar por Python: son la mitad de la hoja.
    _CELDA_CON_DATOS = re.compile(r'<c r="([A-Z]+)([0-9]+)"([^>/]*)>(.*?)</c>', re.S)
    _REFERENCIA = re.compile(r'<c r="([A-Z]+)([0-9]+)"')
    _ATRIBUTO_T = re.compile(r'\bt="([^"]*)"')
    _ATRIBUTO_S = re.compile(r'\bs="([0-9]*)"')

    def _celdas_rapido(self, xml):
        """(celdas, max_fila, max_col, hubo_celdas, resto) leyendo el XML con expresiones
        regulares; None si la hoja no tiene la forma de siempre (va por _celdas_et)."""
        import html
        import xml.etree.ElementTree as ET
        texto = xml.decode('utf-8')
        i = texto.find('<sheetData')
        if i < 0 or texto.count('<c ') != texto.count('<c r="') or '<x:' in texto:
            return None
        if texto.startswith('<sheetData/>', i):
            j = i + len('<sheetData/>')
        else:
            j = texto.find('</sheetData>', i)
            if j < 0:
                return None
            j += len('</sheetData>')
        datos = texto[i:j]
        celdas = {}
        for letras, fila, atributos, contenido in self._CELDA_CON_DATOS.findall(datos):
            if not contenido:
                continue
            t = self._ATRIBUTO_T.search(atributos)
            tipo = t.group(1) if t else 'n'
            if tipo == 'inlineStr':
                if '<is' not in contenido:
                    continue
                nodo = ET.fromstring(f'<c xmlns="{_NS_XLSX[1:-1]}">{contenido}</c>').find(_NS_XLSX + 'is')
                valor = _texto_de(nodo)
            else:
                a = contenido.find('<v>')
                if a < 0:
                    if '<v' in contenido:
                        return None   # <v> con atributos: que lo lea el camino general
                    continue
                b = contenido.find('</v>', a)
                valor = contenido[a + 3:b]
                if not valor:
                    continue
                if '&' in valor or '\r' in valor:
                    valor = html.unescape(valor.replace('\r\n', '\n').replace('\r', '\n'))
                s = self._ATRIBUTO_S.search(atributos)
                valor = self._valor(tipo, valor, int(s.group(1) or 0) if s else 0)
            celdas[(int(fila), _columna_a_numero(letras))] = _Celda(valor)
        referencias = self._REFERENCIA.findall(datos)
        max_fila = max_col = 1
        if referencias:
            max_fila = max(int(f) for _, f in referencias)
            max_col = max(_columna_a_numero(l) for l in {l for l, _ in referencias})
        return celdas, max_fila, max_col, bool(referencias), ET.fromstring(texto[:i] + texto[j:])

    def _celdas_et(self, xml):
        """Lo mismo que _celdas_rapido, recorriendo el árbol XML celda por celda."""
        import xml.etree.ElementTree as ET
        raiz = ET.fromstring(xml)
        celdas = {}
        max_fila = max_col = 1
        hubo_celdas = False
        TAG_ROW, TAG_V, TAG_IS = _NS_XLSX + 'row', _NS_XLSX + 'v', _NS_XLSX + 'is'
        datos = raiz.find(_NS_XLSX + 'sheetData')
        fila_actual = 0
        for fila in (datos if datos is not None else ()):
            if fila.tag != TAG_ROW:
                continue
            r = fila.get('r')
            fila_actual = int(float(r)) if r is not None else fila_actual + 1
            col_actual = 0
            for c in fila:
                ref = c.get('r')
                if ref:
                    letras = ref.rstrip(_DIGITOS)
                    f, col_actual = int(ref[len(letras):]), _columna_a_numero(letras)
                else:
                    f, col_actual = fila_actual, col_actual + 1
                hubo_celdas = True
                max_fila, max_col = max(max_fila, f), max(max_col, col_actual)
                tipo = c.get('t', 'n')
                if tipo == 'inlineStr':
                    hijo = c.find(TAG_IS)
                    valor = _texto_de(hijo) if hijo is not None else None
                else:
                    valor = c.findtext(TAG_V) or None
                    if valor is not None:
                        valor = self._valor(tipo, valor, int(c.get('s') or 0))
                if valor is not None:
                    celdas[(f, col_actual)] = _Celda(valor)
        return celdas, max_fila, max_col, hubo_celdas, raiz

    def _leer(self):
        import xml.etree.ElementTree as ET
        libro = self._libro
        xml = libro._zip.read(self._ruta)
        celdas, max_fila, max_col, hubo_celdas, raiz = self._celdas_rapido(xml) or self._celdas_et(xml)

        # Celdas combinadas: solo la de arriba a la izquierda conserva su valor.
        combinadas = []
        nodo = raiz.find(_NS_XLSX + 'mergeCells')
        for m in (nodo.iter(_NS_XLSX + 'mergeCell') if nodo is not None else ()):
            f1, c1, f2, c2 = _rango(m.get('ref'))
            combinadas.append((f1, c1, f2, c2))
            for f in range(f1, f2 + 1):
                for c in range(c1, c2 + 1):
                    if (f, c) != (f1, c1):
                        celdas[(f, c)] = None   # None = celda combinada (MergedCell)

        def celda(f, c):
            actual = celdas.get((f, c), _CELDA_VACIA)
            if actual is _CELDA_VACIA:
                actual = celdas[(f, c)] = _Celda()
            return actual

        def enlazar(destino, objetivo, lugar):
            # Como el setter de openpyxl: una celda vacía toma el enlace como valor.
            destino.hyperlink = _Enlace(objetivo)
            if destino.value is None:
                destino.value = objetivo or lugar

        rels = libro._rels_de(self._ruta)
        nodo = raiz.find(_NS_XLSX + 'hyperlinks')
        for h in (nodo.iter(_NS_XLSX + 'hyperlink') if nodo is not None else ()):
            id_rel = h.get(_NS_REL_DOC + 'id')
            objetivo, lugar = (rels[id_rel][0] if id_rel else None), h.get('location')
            ref = h.get('ref')
            f1, c1, f2, c2 = _rango(ref)
            if ':' in ref:
                for f in range(f1, f2 + 1):
                    for c in range(c1, c2 + 1):
                        if celdas.get((f, c), 0) is not None:   # las combinadas se saltan
                            enlazar(celda(f, c), objetivo, lugar)
            else:
                if celdas.get((f1, c1), 0) is None:   # cae en una combinada: va a su ancla
                    f1, c1 = next((a, b) for a, b, x, y in combinadas if a <= f1 <= x and b <= c1 <= y)
                enlazar(celda(f1, c1), objetivo, lugar)

        # Los comentarios crean su celda (openpyxl hace ws[ref].comment = ...).
        for destino, tipo in rels.values():
            if tipo.endswith('/comments'):
                for com in ET.fromstring(libro._zip.read(destino)).iter(_NS_XLSX + 'comment'):
                    f, c, _, _ = _rango(com.get('ref'))
                    if celdas.get((f, c), 0) is not None:
                        celda(f, c)

        self._hubo_celdas = hubo_celdas
        self._celdas = celdas
        # Las vacías no se guardan, pero cuentan para el tamaño (openpyxl sí las crea).
        if celdas:
            max_fila = max(max_fila, max(f for f, _ in celdas))
            max_col = max(max_col, max(c for _, c in celdas))
        self._tamano = (max_fila, max_col)

    def iter_rows(self):
        if self._celdas is None:
            self._leer()
        celdas = self._celdas
        if not self._hubo_celdas:
            return iter(())
        max_fila, max_col = self._tamano
        vacia = _CELDA_VACIA

        def filas():
            obtener = celdas.get
            columnas = range(1, max_col + 1)
            for f in range(1, max_fila + 1):
                yield tuple((obtener((f, c)) or vacia) for c in columnas)
        return filas()


class LibroRapido:
    """Lo que buscador.descargar_linea_credito usa de un Workbook de openpyxl:
    .worksheets, y de cada hoja .title e iter_rows() con celdas .value/.hyperlink.target.
    Cada hoja se lee recién cuando se recorre."""

    def __init__(self, fuente):
        import zipfile
        import xml.etree.ElementTree as ET
        from openpyxl.styles.numbers import builtin_format_code, is_date_format, is_timedelta_format
        from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
        if isinstance(fuente, str):   # en memoria: así el archivo no queda bloqueado en Windows
            with open(fuente, 'rb') as f:
                fuente = io.BytesIO(f.read())
        self._zip = zipfile.ZipFile(fuente)
        nombres = set(self._zip.namelist())
        self._cache_rels = {}

        libro = ET.fromstring(self._zip.read('xl/workbook.xml'))
        pr = libro.find(_NS_XLSX + 'workbookPr')
        fecha1904 = pr is not None and pr.get('date1904') in ('1', 'true')
        self._epoca = CALENDAR_MAC_1904 if fecha1904 else CALENDAR_WINDOWS_1900

        rels = self._rels_de('xl/workbook.xml')
        self.worksheets = []
        for hoja in libro.iter(_NS_XLSX + 'sheet'):
            id_rel = hoja.get(_NS_REL_DOC + 'id')
            if not id_rel:
                continue
            ruta, tipo = rels[id_rel]
            if ruta not in nombres or 'chartsheet' in tipo:
                continue
            if not tipo.endswith('/worksheet'):
                raise ValueError(f"hoja de tipo desconocido: {tipo}")
            self.worksheets.append(HojaRapida(self, hoja.get('name'), ruta))

        ruta_textos = next((r for r, t in rels.values() if t.endswith('/sharedStrings')), None)
        self._compartidos = []
        if ruta_textos and ruta_textos in nombres:
            for si in ET.fromstring(self._zip.read(ruta_textos)).iter(_NS_XLSX + 'si'):
                self._compartidos.append(_texto_de(si).replace('x005F_', ''))

        # Qué estilos de celda son de fecha o de duración (como openpyxl _normalise_numbers).
        self._formatos_fecha, self._formatos_duracion = set(), set()
        ruta_estilos = next((r for r, t in rels.values() if t.endswith('/styles')), None)
        if ruta_estilos and ruta_estilos in nombres:
            estilos = ET.fromstring(self._zip.read(ruta_estilos))
            propios = {}
            nodo = estilos.find(_NS_XLSX + 'numFmts')
            for n in (nodo.iter(_NS_XLSX + 'numFmt') if nodo is not None else ()):
                propios[int(n.get('numFmtId'))] = n.get('formatCode')
            nodo = estilos.find(_NS_XLSX + 'cellXfs')
            for i, xf in enumerate(nodo.iter(_NS_XLSX + 'xf') if nodo is not None else ()):
                id_formato = int(xf.get('numFmtId', 0))
                formato = propios[id_formato] if id_formato in propios else builtin_format_code(id_formato)
                if is_date_format(formato):
                    self._formatos_fecha.add(i)
                if is_timedelta_format(formato):
                    self._formatos_duracion.add(i)

    def _rels_de(self, ruta):
        if ruta not in self._cache_rels:
            self._cache_rels[ruta] = _rels(self._zip, _ruta_rels(ruta))
        return self._cache_rels[ruta]


# openpyxl busca el destino de cada hipervínculo recorriendo TODA la lista de relaciones de
# la hoja: con ~20.000 hipervínculos por línea son millones de comparaciones (en Libre
# Inversión Menor, 9 de los 13 s de lectura). Con un índice por Id la búsqueda es inmediata
# y el resultado, el mismo. Sirve para lo que todavía se abre con openpyxl (el respaldo de
# arriba, el visor de archivos Excel de la versión de navegador).
try:
    from openpyxl.packaging.relationship import RelationshipList

    def _relacion_por_id(self, key):
        indice = self.__dict__.get('_axio_indice')
        if indice is None or indice[0] != len(self):   # la lista cambió: se rehace
            mapa = {}
            for r in self:
                mapa.setdefault(r.Id, r)   # como el original: gana la primera con ese Id
            indice = self.__dict__['_axio_indice'] = (len(self), mapa)
        r = indice[1].get(key)
        if r is None:
            raise KeyError("Unknown relationship: {0}".format(key))
        return r

    RelationshipList.get = _relacion_por_id
except ImportError:   # otra versión de openpyxl: queda el método original, solo más lento
    pass


# ==============================================================================
# HILOS Y RECOLECTOR DE BASURA (V10.8)
# ==============================================================================
@contextmanager
def gc_fuera_del_hilo(app=None):
    """Evita que el recolector de basura corra DENTRO de un hilo de trabajo.

    El problema real, visto en producción: una búsqueda de Axio lee decenas de archivos
    Excel; eso reserva y libera mucha memoria, así que el recolector se dispara -- y se
    dispara en ESE hilo. Si en ese momento quedaba basura de tkinter (una fuente, una
    variable, una imagen que la interfaz descartó al redibujar), su __del__ llama a Tcl
    desde un hilo que no es el principal. Primero aparecen mensajes sueltos de
    "main thread is not in main loop" y al final Tcl_AsyncDelete cierra la aplicación
    entera, sin diálogo y sin nada en el log de errores.

    Aquí se apaga el recolector mientras dura el hilo y se vuelve a encender al terminar;
    la recolección se pide en el hilo principal con after(), que es donde sí se puede
    hablar con Tcl. No se pierde memoria: el conteo de referencias sigue liberando lo
    normal, solo se aplazan los ciclos unos segundos.
    """
    estaba_activo = gc.isenabled()
    if estaba_activo:
        gc.disable()
    try:
        yield
    finally:
        if estaba_activo:
            gc.enable()
            if app is not None:
                try:
                    app.after(0, gc.collect)
                except Exception:
                    pass  # la ventana ya se cerró: no hay nada que recolectar
