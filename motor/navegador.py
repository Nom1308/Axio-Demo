"""Adaptador entre la pestaña del navegador y el MotorWeb de Axio.

Corre dentro de Pyodide, en un Web Worker (ver estatico/trabajador.js). Hace lo mismo que
web/app.py en el servidor -- recibe /api/... y devuelve la respuesta -- pero sin Flask,
sin usuarios y sin red: el "servidor" es esta misma pestaña.

Con sesión de Google, cada descarga lleva el token de la persona (ver fijar_token): las
hojas se leen con SUS permisos y pueden ser privadas.

Tres diferencias con el servidor, todas por el entorno:
  - No hay hilos en el navegador: la carga de datos se hace de una vez, al recibir el
    config, en vez de en un hilo aparte.
  - urllib no tiene sockets aquí: urlopen se cambia por una petición del navegador
    (XMLHttpRequest síncrono, permitido dentro de un Worker). Google Sheets responde con
    CORS abierto, así que el navegador puede bajar las hojas directamente.
  - No hay cierres locales: una página web no puede leer carpetas del equipo.

Todo vive en la memoria de la pestaña (/datos es un sistema de archivos en RAM). Al
cerrarla o recargarla no queda nada.
"""

import io
import json
import os
import re
import sys
import email.message
import urllib.error
import urllib.request
from urllib.parse import parse_qs, unquote, urlsplit

CARPETA_DATOS = os.environ.setdefault('AXIO_CARPETA_DATOS', '/datos')
EN_NAVEGADOR = sys.platform == 'emscripten'
USUARIO = 'demo'
LARGO_MAXIMO_BUSQUEDA = 200   # el mismo límite que web/app.py
TIPO_JSON = 'application/json'
TIPO_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


# ------------------------------------------------------------------ red
class _Respuesta:
    """Lo mínimo de http.client.HTTPResponse que usan pandas y axio/nucleo/utils."""

    def __init__(self, url, estado, datos, tipo):
        self.url = url
        self.status = estado
        self._buffer = io.BytesIO(datos)
        self.headers = email.message.Message()
        if tipo:
            self.headers['Content-Type'] = tipo

    def read(self, n=-1):
        return self._buffer.read(n)

    def getcode(self):
        return self.status

    def close(self):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()


# Token de Google de quien inició sesión (ver estatico/puente.js). Con él, cada descarga a
# Google va con los permisos de esa persona y las hojas pueden ser privadas. Vacío: las
# descargas son anónimas y solo funcionan con hojas compartidas por enlace.
_token_google = ''
_HOSTS_GOOGLE = ('docs.google.com', 'sheets.googleapis.com', 'www.googleapis.com')


def fijar_token(token):
    global _token_google
    _token_google = token or ''


# Una exportación grande (Libre Inversión Mayor pesa ~6 MB en xlsx) a veces falla sin
# respuesta: Google corta o devuelve un error sin cabeceras CORS, y el navegador lo ve como
# "sin conexión" (estado 0). Suele salir bien al segundo intento.
_INTENTOS_DESCARGA = 3


def _pedir(url):
    import js
    xhr = js.XMLHttpRequest.new()
    xhr.open('GET', url, False)
    xhr.responseType = 'arraybuffer'
    if _token_google and urlsplit(url).hostname in _HOSTS_GOOGLE:
        xhr.setRequestHeader('Authorization', 'Bearer ' + _token_google)
    try:
        xhr.send()
    except Exception:
        return None
    return xhr if xhr.status != 0 else None


def _urlopen_navegador(url, data=None, timeout=None, **_):
    if isinstance(url, urllib.request.Request):
        url = url.full_url
    xhr = None
    for _intento in range(_INTENTOS_DESCARGA):
        xhr = _pedir(url)
        if xhr is not None and xhr.status < 500 and xhr.status != 429:
            break
    if xhr is None:
        raise urllib.error.URLError(
            f"{urlsplit(url).netloc} no respondió después de {_INTENTOS_DESCARGA} intentos. "
            "Si el resto de las hojas cargó, es momentáneo: prueba «Refrescar datos»")
    import js
    if xhr.status in (401, 403, 404) and 'google' in (urlsplit(url).hostname or ''):
        motivo = ("tu cuenta de Google no tiene permiso para ver esta hoja; pide acceso a quien la administra"
                  if _token_google else
                  "la hoja no está compartida por enlace; inicia sesión con Google o pide acceso")
        raise urllib.error.HTTPError(url, xhr.status, motivo, None, None)
    if xhr.status >= 400:
        raise urllib.error.HTTPError(url, xhr.status, f"HTTP {xhr.status}", None, None)
    datos = bytes(js.Uint8Array.new(xhr.response).to_py())
    return _Respuesta(xhr.responseURL or url, xhr.status, datos, xhr.getResponseHeader('Content-Type'))


if EN_NAVEGADOR:
    # Antes de importar axio: pandas y axio/nucleo/utils llaman a urllib.request.urlopen
    # por atributo, así que ven este reemplazo.
    urllib.request.urlopen = _urlopen_navegador


def _avisar(mensaje):
    if EN_NAVEGADOR and mensaje:
        import js
        js.postMessage(js.JSON.parse(json.dumps({'tipo': 'progreso', 'mensaje': mensaje})))


# ------------------------------------------------------------------ motor
from axio.nucleo.rutas import CONFIG_FILE   # noqa: E402  (después de fijar la carpeta)
from web.servicio import MotorWeb           # noqa: E402


class MotorNavegador(MotorWeb):
    """El MotorWeb del servidor, con el aviso de progreso conectado a la pestaña."""

    @property
    def mensaje_carga(self):
        return self._mensaje_carga

    @mensaje_carga.setter
    def mensaje_carga(self, valor):
        self._mensaje_carga = valor
        _avisar(valor)

    def iniciar_carga(self, forzar=False):
        # Sin hilos en el navegador. La pestaña nunca llega aquí porque solo busca con
        # los datos ya cargados; si llegara, que no intente arrancar un hilo.
        return False


_motor = None


def configurar(texto_config):
    """Recibe el contenido de config_axio.json, lo deja en /datos y descarga todo.
    Devuelve el estado (JSON) para pintar las fuentes."""
    global _motor
    config = json.loads(texto_config)
    if not isinstance(config, dict):
        raise ValueError("El archivo no tiene el formato de config_axio.json.")
    config['buscador_incluir_archivos_locales'] = False
    os.makedirs(CARPETA_DATOS, exist_ok=True)
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(config, f, ensure_ascii=False)
    _motor = MotorNavegador()
    _motor.cargar_ahora(forzar=True)
    return json.dumps(_motor.estado())


def recargar():
    """El «Refrescar datos» de la barra: vuelve a bajar todo con el mismo config."""
    _motor.cargar_ahora(forzar=True)
    return json.dumps(_motor.estado())


def _estado_con_enlaces():
    """El estado del servidor más el enlace de cada hoja, para abrirla desde el menú del
    usuario. En el servidor las URLs no salen al navegador; aquí sí, porque quien usa la
    pestaña ya tiene el config y entra a cada hoja con sus propios permisos."""
    estado = _motor.estado()
    config = _motor.config
    enlaces = {'extractos': config.get('url_base_global'), 'whatsapp': config.get('url_whatsapp')}
    enlaces.update({l.get('clave'): l.get('url') for l in config.get('lineas_credito', [])})
    for f in estado.get('fuentes', []):
        url = str(enlaces.get(f['clave']) or '').strip()
        f['url'] = url if url.startswith('https://') else None
    return estado


# ------------------------------------------------------------------ archivos de Excel en Drive
# Algunas obligaciones apuntan a un .xlsx guardado en Drive, no a una hoja de Google: la API
# de Sheets no los lee ("must not be an Office file"). Se bajan por la API de Drive y se leen
# con openpyxl, devolviendo la MISMA forma que la API de Sheets (rowData, merges...) para que
# estatico/hojas.js los dibuje con el mismo código. Solo lectura.
_cache_excel = {}          # id -> bytes del archivo (cambiar de pestaña no lo vuelve a bajar)
_MAX_CACHE_EXCEL = 6


def _color_excel(color):
    """'FFRRGGBB' de openpyxl -> {red, green, blue} de 0 a 1. Los colores de tema se ignoran."""
    rgb = getattr(color, 'rgb', None) if color is not None else None
    if not isinstance(rgb, str) or len(rgb) != 8 or getattr(color, 'type', 'rgb') != 'rgb':
        return None
    try:
        r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (2, 4, 6))
    except ValueError:
        return None
    return {'red': r, 'green': g, 'blue': b}


def _miles(n, decimales=0):
    texto = f"{abs(n):,.{decimales}f}".replace(',', '_').replace('.', ',').replace('_', '.')
    return ('-' if n < 0 else '') + texto


def _texto_excel(valor, formato):
    """El valor como se vería en Excel, en formato colombiano (1.200.000 y 05/10/2026)."""
    from datetime import date, datetime, time
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.strftime('%d/%m/%Y') if not (valor.hour or valor.minute) else valor.strftime('%d/%m/%Y %H:%M')
    if isinstance(valor, date):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, time):
        return valor.strftime('%H:%M')
    if isinstance(valor, bool):
        return 'VERDADERO' if valor else 'FALSO'
    if isinstance(valor, (int, float)):
        formato = formato or 'General'
        if '%' in formato:
            return _miles(valor * 100, 2 if '.0' in formato else 0) + '%'
        decimales = len(formato.split('.')[1].split(';')[0].rstrip('_)" ')) if '.0' in formato else 0
        if formato == 'General':
            decimales = 0 if float(valor).is_integer() else min(4, len(repr(float(valor)).split('.')[1]))
            texto = _miles(valor, decimales) if abs(valor) >= 1000 else (f"{valor:.{decimales}f}".replace('.', ','))
        else:
            texto = _miles(valor, decimales) if ('#,##' in formato or '$' in formato) else f"{valor:.{decimales}f}".replace('.', ',')
        return ('$ ' + texto) if '$' in formato else texto
    return str(valor)


def leer_excel(id_archivo, nombre_hoja, max_filas, max_columnas):
    """JSON con {hojas: [...títulos], hoja, libro: <forma de la API de Sheets>}."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    if id_archivo not in _cache_excel:
        url = f"https://www.googleapis.com/drive/v3/files/{id_archivo}?alt=media&supportsAllDrives=true"
        with urllib.request.urlopen(url) as resp:
            _cache_excel[id_archivo] = resp.read()
        while len(_cache_excel) > _MAX_CACHE_EXCEL:
            _cache_excel.pop(next(iter(_cache_excel)))
    wb = load_workbook(io.BytesIO(_cache_excel[id_archivo]), data_only=True)
    visibles = [ws for ws in wb.worksheets if ws.sheet_state == 'visible'] or wb.worksheets
    ws = next((w for w in visibles if w.title == nombre_hoja), visibles[0])
    filas = min(ws.max_row or 1, max_filas)
    columnas = min(ws.max_column or 1, max_columnas)

    alineaciones = {'left': 'LEFT', 'center': 'CENTER', 'centerContinuous': 'CENTER', 'right': 'RIGHT'}
    datos_filas = []
    for fila in ws.iter_rows(min_row=1, max_row=filas, max_col=columnas):
        valores = []
        for c in fila:
            formato = {}
            relleno = c.fill
            if relleno is not None and relleno.fill_type == 'solid':
                fondo = _color_excel(relleno.fgColor)
                if fondo:
                    formato['backgroundColor'] = fondo
            fuente = c.font
            texto = {}
            if fuente is not None:
                if fuente.b:
                    texto['bold'] = True
                if fuente.i:
                    texto['italic'] = True
                letra = _color_excel(fuente.color)
                if letra:
                    texto['foregroundColor'] = letra
            if texto:
                formato['textFormat'] = texto
            horizontal = c.alignment.horizontal if c.alignment is not None else None
            if horizontal in alineaciones:
                formato['horizontalAlignment'] = alineaciones[horizontal]
            celda = {'formattedValue': _texto_excel(c.value, c.number_format), 'effectiveFormat': formato}
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                celda['effectiveValue'] = {'numberValue': c.value}
            if c.hyperlink is not None and c.hyperlink.target:
                celda['hyperlink'] = c.hyperlink.target
            valores.append(celda)
        datos_filas.append({'values': valores})

    anchos = []
    for i in range(1, columnas + 1):
        dim = ws.column_dimensions.get(get_column_letter(i))
        ancho = dim.width if dim is not None and dim.width else 8.43
        anchos.append({'pixelSize': round(ancho * 7 + 5)})
    altos = []
    for i in range(1, filas + 1):
        dim = ws.row_dimensions.get(i)
        altos.append({'pixelSize': round(dim.height * 4 / 3) if dim is not None and dim.height else 21})
    merges = [{'startRowIndex': r.min_row - 1, 'endRowIndex': r.max_row,
               'startColumnIndex': r.min_col - 1, 'endColumnIndex': r.max_col}
              for r in ws.merged_cells.ranges if r.min_row <= filas and r.min_col <= columnas]
    return json.dumps({
        'hojas': [w.title for w in visibles],
        'hoja': ws.title,
        'recortada': (ws.max_row or 0) > max_filas or (ws.max_column or 0) > max_columnas,
        'libro': {'sheets': [{'merges': merges, 'data': [{
            'columnMetadata': anchos, 'rowMetadata': altos, 'rowData': datos_filas}]}]},
    }, ensure_ascii=False)


def _json(datos, estado=200):
    return estado, TIPO_JSON, json.dumps(datos, ensure_ascii=False).encode('utf-8')


def atender(url):
    """Responde una llamada GET de app.js igual que web/app.py.
    Devuelve (código HTTP, tipo de contenido, cuerpo en bytes)."""
    if _motor is None:
        return _json({'error': "Primero carga el archivo config_axio.json."}, 503)
    partes = urlsplit(url)
    ruta = [unquote(p) for p in partes.path.strip('/').split('/')]

    if ruta == ['api', 'estado']:
        return _json(_estado_con_enlaces())

    if ruta == ['api', 'buscar']:
        termino = re.sub(r'\s+', ' ', parse_qs(partes.query).get('q', [''])[0]).strip()
        if not termino:
            return _json({'error': "Escribe algo para buscar."}, 400)
        if len(termino) > LARGO_MAXIMO_BUSQUEDA:
            return _json({'error': f"La búsqueda no puede pasar de {LARGO_MAXIMO_BUSQUEDA} caracteres."}, 400)
        resultado = _motor.buscar(termino, USUARIO)
        if resultado is None:
            return _json({'cargando': True, 'mensaje': "Cargando los datos..."}, 202)
        return _json(resultado)

    if len(ruta) == 4 and ruta[:2] == ['api', 'detalle'] and ruta[3].isdigit():
        detalle = _motor.detalle(ruta[2], int(ruta[3]), USUARIO)
        if detalle is None:
            return _json({'error': "Ese resultado ya no está disponible. Vuelve a buscar."}, 404)
        return _json(detalle)

    if len(ruta) == 3 and ruta[:2] == ['api', 'exportar']:
        exportado = _motor.exportar(ruta[2], USUARIO)
        if exportado is None:
            return _json({'error': "Esa búsqueda ya no está disponible. Vuelve a buscar."}, 404)
        return 200, TIPO_XLSX, exportado[0]

    return _json({'error': "No encontrado."}, 404)
